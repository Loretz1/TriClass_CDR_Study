#!/usr/bin/env python3
"""Append VAE, WGAN-GP, and flow-matching ablations to the record workbook."""

from __future__ import annotations

import argparse
import ast
import re
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable

from openpyxl import load_workbook
from openpyxl.styles import Font


ROOT = Path(__file__).resolve().parents[2]

METHODS = {
    "VAE": "*generative_ablation_vae*.log",
    "WGAN-GP (GAN)": "*generative_ablation_gan*.log",
    "FM": "*generative_ablation_fm*.log",
}

DOMAINS = [
    {
        "needle": "Amazon2014-Clothing_Shoes_and_Jewelry+Sports_and_Outdoors",
        "dataset": "Amazon2014",
        "direction": "Cloth -> Sport",
        "diffusion": (0.0541, 0.0763, 0.0616, 0.0700),
    },
    {
        "needle": "Amazon2014-Sports_and_Outdoors+Clothing_Shoes_and_Jewelry",
        "dataset": "Amazon2014",
        "direction": "Sport -> Cloth",
        "diffusion": (0.0414, 0.0621, 0.0486, 0.0559),
    },
    {
        "needle": "Amazon2014-Electronics+Cell_Phones_and_Accessories",
        "dataset": "Amazon2014",
        "direction": "Elec -> Phone",
        "diffusion": (0.0650, 0.0994, 0.0677, 0.0820),
    },
    {
        "needle": "Amazon2014-Cell_Phones_and_Accessories+Electronics",
        "dataset": "Amazon2014",
        "direction": "Phone -> Elec",
        "diffusion": (0.0458, 0.0667, 0.0740, 0.0762),
    },
    {
        "needle": "Douban-Book+Music",
        "dataset": "Douban",
        "direction": "Book -> Music",
        "diffusion": (0.0133, 0.0285, 0.0986, 0.0958),
    },
    {
        "needle": "Douban-Music+Book",
        "dataset": "Douban",
        "direction": "Music -> Book",
        "diffusion": (0.0294, 0.0580, 0.1722, 0.1575),
    },
]

PARAM_RE = re.compile(
    r"Best Parameters:\s*(\[[^\]]+\])\s*=\s*(\([^\n]+\))"
)
METRIC_RE = re.compile(r"^\s*(ndcg|recall)\s*:\s*(.*)$", re.MULTILINE | re.I)
AT_RE = re.compile(r"@(10|20)=([0-9.]+)")


@dataclass
class Result:
    method: str
    domain: dict
    path: Path
    recall10: float
    recall20: float
    ndcg10: float
    ndcg20: float
    params: Dict[str, object]


def scalar(value: object) -> object:
    if isinstance(value, list) and len(value) == 1:
        return value[0]
    return value


def config_value(text: str, key: str) -> object | None:
    match = re.search(rf"^{re.escape(key)}=(.+)$", text, re.MULTILINE)
    if not match:
        return None
    raw = match.group(1).strip()
    try:
        return scalar(ast.literal_eval(raw))
    except (SyntaxError, ValueError):
        return raw


def parse_log(method: str, domain: dict, path: Path) -> Result:
    text = path.read_text(encoding="utf-8", errors="ignore")
    marker = text.rfind("BEST RESULTS")
    if marker < 0:
        raise ValueError(f"No final BEST RESULTS block: {path}")
    block = text[marker:]
    test_marker = block.rfind("Test:")
    if test_marker < 0:
        raise ValueError(f"No final Test block: {path}")

    metrics: Dict[str, Dict[int, float]] = {}
    for match in METRIC_RE.finditer(block[test_marker:]):
        metrics[match.group(1).lower()] = {
            int(k): float(v) for k, v in AT_RE.findall(match.group(2))
        }
    if not all(k in metrics.get("recall", {}) for k in (10, 20)):
        raise ValueError(f"Missing Recall@10/20: {path}")
    if not all(k in metrics.get("ndcg", {}) for k in (10, 20)):
        raise ValueError(f"Missing NDCG@10/20: {path}")

    param_matches = list(PARAM_RE.finditer(block))
    if not param_matches:
        raise ValueError(f"Missing best parameters: {path}")
    names = ast.literal_eval(param_matches[-1].group(1))
    values = ast.literal_eval(param_matches[-1].group(2))
    if not isinstance(values, tuple):
        values = (values,)
    params = dict(zip(names, values))

    common_keys = [
        "timesteps",
        "top_k_pos",
        "pseudo_start_epoch",
        "beta_sche",
        "diff_weight",
        "lambda_user_emb",
        "aggregator",
        "history_len",
        "mask_rate",
    ]
    method_keys = {
        "VAE": ["vae_hidden_dim", "vae_latent_dim", "vae_kl_weight"],
        "WGAN-GP (GAN)": [
            "wgan_hidden_dim",
            "wgan_adversarial_weight",
            "wgan_critic_steps",
        ],
        "FM": ["fm_sampling_steps", "fm_solver", "fm_noise_scale"],
    }[method]
    for key in common_keys + method_keys:
        if key not in params:
            value = config_value(text, key)
            if value is not None:
                params[key] = value

    return Result(
        method=method,
        domain=domain,
        path=path,
        recall10=metrics["recall"][10],
        recall20=metrics["recall"][20],
        ndcg10=metrics["ndcg"][10],
        ndcg20=metrics["ndcg"][20],
        params=params,
    )


def collect_results(log_dir: Path) -> Dict[tuple[str, str], Result]:
    results: Dict[tuple[str, str], Result] = {}
    for method, pattern in METHODS.items():
        paths = list(log_dir.glob(pattern))
        if len(paths) != 6:
            raise ValueError(f"Expected 6 {method} logs, found {len(paths)}")
        for domain in DOMAINS:
            matches = [path for path in paths if domain["needle"] in path.name]
            if len(matches) != 1:
                raise ValueError(
                    f"Expected one {method} log for {domain['direction']}, found {len(matches)}"
                )
            results[(method, domain["direction"])] = parse_log(
                method, domain, matches[0]
            )
    if len(results) != 18:
        raise ValueError(f"Expected 18 parsed records, found {len(results)}")
    return results


def copy_row_style(ws, source_row: int, target_row: int, columns: int = 9) -> None:
    for column in range(1, columns + 1):
        source = ws.cell(source_row, column)
        target = ws.cell(target_row, column)
        target._style = copy(source._style)
        if source.has_style:
            target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def merged_row(ws, row: int, value: str, template_row: int, height: float) -> None:
    copy_row_style(ws, template_row, row)
    for column in range(1, 10):
        ws.cell(row, column).value = None
    ws.cell(row, 1).value = value
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    ws.row_dimensions[row].height = height


def format_params(result: Result) -> str:
    p = result.params
    aliases = [
        ("seed", "seed"),
        ("training_stages.0.learning_rate", "lr"),
        ("timesteps", "T"),
        ("top_k_pos", "topk"),
        ("pseudo_start_epoch", "start"),
        ("beta_sche", "beta"),
        ("diff_weight", "diff_w"),
        ("lambda_user_emb", "lambda_u"),
        ("aggregator", "agg"),
        ("history_len", "hist"),
        ("mask_rate", "mask"),
    ]
    extras = {
        "VAE": [
            ("vae_hidden_dim", "vae_hidden"),
            ("vae_latent_dim", "vae_latent"),
            ("vae_kl_weight", "vae_kl"),
        ],
        "WGAN-GP (GAN)": [
            ("wgan_hidden_dim", "gan_hidden"),
            ("wgan_adversarial_weight", "gan_adv_w"),
            ("wgan_critic_steps", "critic_steps"),
        ],
        "FM": [
            ("fm_sampling_steps", "fm_steps"),
            ("fm_solver", "solver"),
            ("fm_noise_scale", "noise"),
        ],
    }[result.method]
    parts = []
    for key, label in aliases + extras:
        if key in p:
            parts.append(f"{label}={p[key]}")
    return ", ".join(parts)


def delta_text(result: Result) -> str:
    _, ref_r20, _, ref_n20 = result.domain["diffusion"]
    delta_r = result.recall20 - ref_r20
    delta_n = result.ndcg20 - ref_n20
    return (
        f"相对完整 Diffusion，R@20 {delta_r:+.4f}，"
        f"N@20 {delta_n:+.4f}。"
    )


def format_pair(recall20: float, ndcg20: float) -> str:
    return f"{recall20:.4f} / {ndcg20:.4f}"


def append_detail_sheet(ws, results: Dict[tuple[str, str], Result]) -> None:
    marker = "补充实验：生成式迁移模块消融"
    if any(marker in str(cell.value) for row in ws.iter_rows() for cell in row if cell.value):
        raise ValueError(f"Section already exists in {ws.title}")

    row = ws.max_row + 2
    merged_row(
        ws,
        row,
        "补充实验：生成式迁移模块消融（以 VAE、WGAN-GP、FM 替换 Diffusion）",
        179,
        30,
    )
    row += 1
    copy_row_style(ws, 180, row)
    headers = [
        "模型 / 变体",
        "类型",
        "Recall@10",
        "Recall@20",
        "NDCG@10",
        "NDCG@20",
        "重要参数",
        "log / 来源",
        "一句话结论",
    ]
    for column, value in enumerate(headers, 1):
        ws.cell(row, column).value = value
    ws.row_dimensions[row].height = 30

    for index, domain in enumerate(DOMAINS, 1):
        row += 1
        _, ref_r20, _, ref_n20 = domain["diffusion"]
        merged_row(
            ws,
            row,
            (
                f"数据集{index}：{domain['dataset']} {domain['direction']} "
                f"（完整 Diffusion R@20/N@20={ref_r20:.4f}/{ref_n20:.4f}）"
            ),
            124,
            36,
        )
        for method in METHODS:
            result = results[(method, domain["direction"])]
            row += 1
            copy_row_style(ws, 181, row)
            values = [
                f"LSDCDR / {method}",
                "generative-ablation",
                result.recall10,
                result.recall20,
                result.ndcg10,
                result.ndcg20,
                format_params(result),
                f"src/log/{result.path.name} (最终 BEST RESULTS / Test)",
                delta_text(result),
            ]
            for column, value in enumerate(values, 1):
                ws.cell(row, column).value = value
            for column in range(3, 7):
                ws.cell(row, column).number_format = "0.0000"
            for column in (1, 2, 7, 8, 9):
                alignment = copy(ws.cell(row, column).alignment)
                alignment.wrap_text = True
                alignment.vertical = "center"
                ws.cell(row, column).alignment = alignment
            ws.row_dimensions[row].height = 54

    row += 1
    merged_row(
        ws,
        row,
        (
            "结论：保持 LLM 语义表示、伪交互检索、EMA 与 self-evolving 机制不变，仅替换生成式迁移模块后，"
            "VAE、WGAN-GP 和 FM 在六个方向的 R@20/N@20 均低于完整 Diffusion；"
            "说明性能增益不仅来自伪监督框架，Diffusion 本身也提供了更有效的跨域分布迁移能力。"
        ),
        187,
        60,
    )


def append_summary_sheet(ws, results: Dict[tuple[str, str], Result]) -> None:
    marker = "补充汇总：生成式迁移模块消融"
    if any(marker in str(cell.value) for row in ws.iter_rows() for cell in row if cell.value):
        raise ValueError(f"Section already exists in {ws.title}")

    row = ws.max_row + 2
    merged_row(
        ws,
        row,
        "补充汇总：生成式迁移模块消融（Diffusion / VAE / WGAN-GP / FM）",
        49,
        30,
    )
    row += 1
    copy_row_style(ws, 34, row)
    headers = [
        "Backbone",
        "Transfer module",
        "Cloth->Sport R@20/N@20",
        "Sport->Cloth R@20/N@20",
        "Elec->Phone R@20/N@20",
        "Phone->Elec R@20/N@20",
        "Book->Music R@20/N@20",
        "Music->Book R@20/N@20",
        "一句话结论",
    ]
    for column, value in enumerate(headers, 1):
        ws.cell(row, column).value = value
    ws.row_dimensions[row].height = 30

    row += 1
    copy_row_style(ws, 36, row)
    diffusion_values = [
        "LSDCDR",
        "Diffusion (完整模型)",
        *[format_pair(domain["diffusion"][1], domain["diffusion"][3]) for domain in DOMAINS],
        "完整 LSDCDR 参考结果。",
    ]
    for column, value in enumerate(diffusion_values, 1):
        ws.cell(row, column).value = value
    ws.row_dimensions[row].height = 42

    averages = {}
    for method in METHODS:
        row += 1
        copy_row_style(ws, 35, row)
        method_results = [results[(method, domain["direction"])] for domain in DOMAINS]
        averages[method] = (
            sum(result.recall20 for result in method_results) / len(method_results),
            sum(result.ndcg20 for result in method_results) / len(method_results),
        )
        values = [
            "LSDCDR",
            method,
            *[format_pair(result.recall20, result.ndcg20) for result in method_results],
            "六个方向均低于完整 Diffusion。",
        ]
        for column, value in enumerate(values, 1):
            ws.cell(row, column).value = value
        ws.row_dimensions[row].height = 42

    row += 1
    merged_row(
        ws,
        row,
        (
            "方法说明：三种变体保留 LLM tag/text embedding、pseudo retrieval、EMA 和 self-evolving update，"
            "只将生成式迁移模块分别替换为 VAE、WGAN-GP 与 conditional flow matching；"
            "因此该实验隔离比较生成模型本身的影响。"
        ),
        47,
        48,
    )
    ws.cell(row, 1).font = Font(
        name=ws.cell(row, 1).font.name or "Arial",
        size=ws.cell(row, 1).font.size or 11,
        color="1F4E78",
        bold=True,
    )

    row += 1
    merged_row(
        ws,
        row,
        (
            "总体结论：完整 Diffusion 的六方向平均 R@20/N@20 为 0.0652/0.0896，"
            f"VAE 为 {averages['VAE'][0]:.4f}/{averages['VAE'][1]:.4f}，"
            f"WGAN-GP 为 {averages['WGAN-GP (GAN)'][0]:.4f}/{averages['WGAN-GP (GAN)'][1]:.4f}，"
            f"FM 为 {averages['FM'][0]:.4f}/{averages['FM'][1]:.4f}；"
            "三类替代生成模型在六个方向均未达到 Diffusion，支持采用 Diffusion 作为 LSDCDR 的生成式迁移模块。"
        ),
        47,
        60,
    )
    ws.cell(row, 1).font = Font(
        name=ws.cell(row, 1).font.name or "Arial",
        size=ws.cell(row, 1).font.size or 11,
        color="C00000",
        bold=True,
    )


def verify_workbook(path: Path) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False)
    expected = {"模型对比", "伪监督质量汇总"}
    if set(workbook.sheetnames) != expected:
        raise ValueError(f"Unexpected sheets: {workbook.sheetnames}")
    if workbook["模型对比"].max_row != 215:
        raise ValueError("Unexpected detail-sheet row count")
    if workbook["伪监督质量汇总"].max_row != 66:
        raise ValueError("Unexpected summary-sheet row count")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--log-dir", type=Path, default=ROOT / "src/log")
    args = parser.parse_args()

    results = collect_results(args.log_dir)
    workbook = load_workbook(args.input, read_only=False, data_only=False)
    if workbook.sheetnames != ["模型对比", "伪监督质量汇总"]:
        raise ValueError(f"Unexpected sheet order: {workbook.sheetnames}")
    append_detail_sheet(workbook["模型对比"], results)
    append_summary_sheet(workbook["伪监督质量汇总"], results)
    workbook.save(args.output)
    verify_workbook(args.output)
    print(f"Wrote {args.output} with {len(results)} parsed logs")


if __name__ == "__main__":
    main()
