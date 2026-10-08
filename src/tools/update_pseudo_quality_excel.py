#!/usr/bin/env python3
"""Build an augmented pseudo-quality experiment workbook from log files.

The sheet follows the CDRec experiment-record template style:
dataset title -> header -> LSDCDR block -> UniCDR block -> fixed pseudo_w=0.1 summary.

For completed logs, the parser uses the final hyper-parameter summary. For
unfinished logs, it falls back to the best already observed Test Summary.
"""

from __future__ import annotations

import argparse
import ast
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]

DOMAINS = [
    ("Clothing_Shoes_and_Jewelry+Sports_and_Outdoors", "数据集1：Amazon2014 Cloth -> Sport"),
    ("Sports_and_Outdoors+Clothing_Shoes_and_Jewelry", "数据集2：Amazon2014 Sport -> Cloth"),
    ("Cell_Phones_and_Accessories+Electronics", "数据集3：Amazon2014 Phone -> Elec"),
]

BACKBONES = ["LSDCDR", "UniCDR"]
VARIANTS = [
    ("OverlapOnly", "Overlap-only"),
    ("ExactHit", "Exact-hit"),
    ("PreferenceNeighbor", "Preference-neighbor"),
    ("Random", "Random"),
    ("Inconsistent", "Inconsistent"),
]
EXTRA_W0_VARIANTS = {"PreferenceNeighbor", "Random", "Inconsistent"}
REFERENCE_FILL = PatternFill("solid", fgColor="D9EAD3")


PARAM_START_RE = re.compile(r"Parameters:(\[[^\]]+\])=\((.*?)\)=+")
SUMMARY_PARAM_RE = re.compile(r"INFO\s+Parameters:\s*(\[[^\]]+\])\s*=\s*\((.*?)\)")
METRIC_RE = re.compile(r"(recall|ndcg)\s*:\s*(.*)", re.I)
AT_RE = re.compile(r"@(\d+)=([0-9.]+)")


@dataclass
class Record:
    path: Path
    backbone: str
    variant: str
    domain_key: str
    params: Dict[str, object]
    recall10: float
    recall20: float
    ndcg10: float
    ndcg20: float
    source: str

    @property
    def pseudo_w(self) -> Optional[float]:
        value = self.params.get("pseudo_rec_weight")
        if value is None:
            return None
        try:
            return float(value)
        except (TypeError, ValueError):
            return None


def parse_literal_tuple(names_text: str, values_text: str) -> Dict[str, object]:
    names = ast.literal_eval(names_text)
    values = ast.literal_eval("(" + values_text + ")")
    if not isinstance(values, tuple):
        values = (values,)
    return dict(zip(names, values))


def parse_metric_line(line: str) -> Optional[Tuple[str, Dict[int, float]]]:
    match = METRIC_RE.search(line)
    if not match:
        return None
    metric_name = match.group(1).lower()
    values = {int(k): float(v) for k, v in AT_RE.findall(match.group(2))}
    return metric_name, values


def infer_meta(path: Path) -> Optional[Tuple[str, str, str]]:
    text = str(path)
    if "log_lsdcdr_pseudo_quality" in text:
        backbone = "LSDCDR"
    elif "log_unicdr_pseudo_quality" in text:
        backbone = "UniCDR"
    else:
        return None

    variant = None
    for key, _ in VARIANTS:
        if f"Pseudo_Quality_{key}" in path.name:
            variant = key
            break
    if variant is None:
        return None

    domain_key = None
    for key, _ in DOMAINS:
        if f"Amazon2014-{key}-" in path.name:
            domain_key = key
            break
    if domain_key is None:
        return None
    return backbone, variant, domain_key


def make_record(
    path: Path,
    meta: Tuple[str, str, str],
    params: Dict[str, object],
    metrics: Dict[str, Dict[int, float]],
    source: str,
) -> Optional[Record]:
    recall = metrics.get("recall", {})
    ndcg = metrics.get("ndcg", {})
    required = [recall.get(10), recall.get(20), ndcg.get(10), ndcg.get(20)]
    if any(v is None for v in required):
        return None
    backbone, variant, domain_key = meta
    return Record(
        path=path,
        backbone=backbone,
        variant=variant,
        domain_key=domain_key,
        params=dict(params),
        recall10=recall[10],
        recall20=recall[20],
        ndcg10=ndcg[10],
        ndcg20=ndcg[20],
        source=source,
    )


def parse_log(path: Path) -> List[Record]:
    meta = infer_meta(path)
    if meta is None:
        return []

    summary_records: List[Record] = []
    epoch_records: List[Record] = []

    current_start_params: Optional[Dict[str, object]] = None
    current_summary_params: Optional[Dict[str, object]] = None
    capture_source: Optional[str] = None
    capture_params: Optional[Dict[str, object]] = None
    capture_metrics: Dict[str, Dict[int, float]] = {}

    def flush_capture() -> None:
        nonlocal capture_source, capture_params, capture_metrics
        if capture_source and capture_params:
            record = make_record(path, meta, capture_params, capture_metrics, capture_source)
            if record:
                if capture_source == "summary":
                    summary_records.append(record)
                else:
                    epoch_records.append(record)
        capture_source = None
        capture_params = None
        capture_metrics = {}

    with path.open("r", encoding="utf-8", errors="ignore") as f:
        for raw_line in f:
            line = raw_line.strip()

            match = PARAM_START_RE.search(line)
            if match:
                flush_capture()
                try:
                    current_start_params = parse_literal_tuple(match.group(1), match.group(2))
                except Exception:
                    current_start_params = None
                continue

            match = SUMMARY_PARAM_RE.search(line)
            if match:
                flush_capture()
                try:
                    current_summary_params = parse_literal_tuple(match.group(1), match.group(2))
                except Exception:
                    current_summary_params = None
                continue

            if "Test Summary" in line and current_start_params:
                flush_capture()
                capture_source = "epoch"
                capture_params = current_start_params
                capture_metrics = {}
                continue

            if line.endswith("INFO    Test:") and current_summary_params:
                flush_capture()
                capture_source = "summary"
                capture_params = current_summary_params
                capture_metrics = {}
                continue

            if capture_source:
                metric = parse_metric_line(line)
                if metric:
                    name, values = metric
                    capture_metrics[name] = values
                    if "recall" in capture_metrics and "ndcg" in capture_metrics:
                        flush_capture()

    flush_capture()
    return summary_records if summary_records else epoch_records


def iter_log_paths() -> Iterable[Path]:
    for log_dir in [ROOT / "src/log_lsdcdr_pseudo_quality", ROOT / "src/log_unicdr_pseudo_quality"]:
        if not log_dir.exists():
            continue
        for path in log_dir.glob("*.log"):
            if infer_meta(path):
                yield path


def best_record(records: Sequence[Record]) -> Optional[Record]:
    if not records:
        return None
    return max(records, key=lambda r: (r.recall20, r.ndcg20, r.recall10, r.ndcg10))


def normalize_param_value(value: object) -> object:
    if isinstance(value, float):
        return round(value, 12)
    if isinstance(value, list):
        return tuple(normalize_param_value(v) for v in value)
    if isinstance(value, tuple):
        return tuple(normalize_param_value(v) for v in value)
    return value


def param_signature(record: Record, exclude: Sequence[str] = ("pseudo_rec_weight",)) -> Tuple[Tuple[str, object], ...]:
    excluded = set(exclude)
    return tuple(
        sorted(
            (key, normalize_param_value(value))
            for key, value in record.params.items()
            if key not in excluded
        )
    )


def paired_weight_record(
    records: Sequence[Record],
    template: Optional[Record],
    pseudo_w: float,
) -> Optional[Record]:
    if template is None:
        return None
    signature = param_signature(template)
    candidates = [
        record
        for record in records
        if record.domain_key == template.domain_key
        and record.backbone == template.backbone
        and record.variant == template.variant
        and record.pseudo_w is not None
        and abs(record.pseudo_w - pseudo_w) < 1e-12
        and param_signature(record) == signature
    ]
    return best_record(candidates)


def best_common_fixed_weight_records(
    records: Sequence[Record],
    domain_key: str,
    backbone: str,
    pseudo_w: float,
) -> Dict[str, Record]:
    """Select one shared non-weight hyper-parameter setting across variants."""
    by_variant: Dict[str, Dict[Tuple[Tuple[str, object], ...], List[Record]]] = {}
    for variant, _ in VARIANTS:
        grouped: Dict[Tuple[Tuple[str, object], ...], List[Record]] = {}
        for record in filter_records(records, domain_key, backbone, variant, pseudo_w=pseudo_w):
            grouped.setdefault(param_signature(record), []).append(record)
        by_variant[variant] = grouped

    common_signatures: Optional[set] = None
    for variant, _ in VARIANTS:
        signatures = set(by_variant[variant].keys())
        common_signatures = signatures if common_signatures is None else common_signatures & signatures

    if not common_signatures:
        return {}

    best_map: Dict[str, Record] = {}
    best_score: Optional[Tuple[float, float, float, float]] = None
    for signature in common_signatures:
        picked = {
            variant: best_record(by_variant[variant][signature])
            for variant, _ in VARIANTS
        }
        if any(record is None for record in picked.values()):
            continue
        values = list(picked.values())
        score = (
            sum(r.recall20 for r in values) / len(values),
            sum(r.ndcg20 for r in values) / len(values),
            sum(r.recall10 for r in values) / len(values),
            sum(r.ndcg10 for r in values) / len(values),
        )
        if best_score is None or score > best_score:
            best_score = score
            best_map = {variant: record for variant, record in picked.items() if record is not None}
    return best_map


def filter_records(
    records: Sequence[Record],
    domain_key: str,
    backbone: str,
    variant: str,
    pseudo_w: Optional[float] = None,
) -> List[Record]:
    out = [
        r
        for r in records
        if r.domain_key == domain_key and r.backbone == backbone and r.variant == variant
    ]
    if pseudo_w is not None:
        out = [r for r in out if r.pseudo_w is not None and abs(r.pseudo_w - pseudo_w) < 1e-12]
    return out


def fmt_metric(value: Optional[float]) -> str:
    if value is None:
        return ""
    return f"{value:.4f}"


def fmt_param_value(value: object) -> str:
    if isinstance(value, float):
        return f"{value:g}"
    return str(value)


def param_summary(record: Optional[Record]) -> str:
    if record is None:
        return ""
    keys = [
        ("training_stages.0.learning_rate", "lr"),
        ("timesteps", "T"),
        ("top_k_pos", "topk"),
        ("pseudo_start_epoch", "start"),
        ("beta_sche", "beta"),
        ("diff_weight", "diff_w"),
        ("lambda_user_emb", "lambda_u"),
        ("history_len", "hist"),
        ("mask_rate", "mask"),
        ("lambda_loss", "lambda"),
        ("pseudo_num_per_batch", "pseudo_bs"),
        ("pseudo_rec_weight", "pseudo_w"),
    ]
    parts = []
    for key, label in keys:
        if key in record.params:
            parts.append(f"{label}={fmt_param_value(record.params[key])}")
    return ", ".join(parts)


def conclusion(backbone: str, variant: str, record: Optional[Record], fixed: Optional[str] = None) -> str:
    if record is None:
        return "未搜索该设置。"
    prefix = f"{fixed}；" if fixed else ""
    mapping = {
        "OverlapOnly": "无冷用户伪监督基线，用作下界。",
        "ExactHit": "真实目标交互上界，显著高于其它设置。",
        "PreferenceNeighbor": "未命中但偏好接近，检验弱质量伪监督价值。",
        "Random": "随机伪监督，检验无偏好信号时的影响。",
        "Inconsistent": "远邻/不一致伪监督，检验错误监督风险。",
    }
    return prefix + mapping.get(variant, "")


def short_variant(backbone: str, variant: str, pseudo_w0: bool = False) -> str:
    label = dict(VARIANTS)[variant]
    if pseudo_w0:
        return f"{backbone} / {label} / pseudo_w=0"
    return f"{backbone} / {label}"


def apply_base_table_styles(ws) -> None:
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = {
        "A": 28,
        "B": 16,
        "C": 12,
        "D": 12,
        "E": 12,
        "F": 12,
        "G": 46,
        "H": 58,
        "I": 56,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in range(1, ws.max_row + 1):
        ws.row_dimensions[row].height = 28


def color_group_best(ws, rows: Sequence[int]) -> None:
    # Color metric cells only: best red, second best yellow, independently per metric.
    best_font = Font(name="SimSun", size=10, bold=True, color="FF0000")
    second_font = Font(name="SimSun", size=10, bold=True, color="FFC000")
    normal_font = Font(name="SimSun", size=10)
    for col in range(3, 7):
        values = []
        for row in rows:
            value = ws.cell(row=row, column=col).value
            try:
                values.append((float(value), row))
            except (TypeError, ValueError):
                pass
        ranked = sorted(values, key=lambda x: x[0], reverse=True)
        for _, row in ranked:
            ws.cell(row=row, column=col).font = normal_font
        if ranked:
            ws.cell(row=ranked[0][1], column=col).font = best_font
        if len(ranked) > 1:
            ws.cell(row=ranked[1][1], column=col).font = second_font


def write_header(ws, row: int, title: str) -> int:
    yellow = PatternFill("solid", fgColor="FFF200")
    grey = PatternFill("solid", fgColor="D9EAD3")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    c = ws.cell(row=row, column=1, value=title)
    c.fill = yellow
    c.font = Font(name="SimSun", size=11, bold=True)
    c.alignment = Alignment(horizontal="left", vertical="center")
    row += 1
    headers = ["模型 / 变体", "类型", "Recall@10", "Recall@20", "NDCG@10", "NDCG@20", "重要参数", "log / 来源", "一句话结论"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.fill = grey
        cell.font = Font(name="SimSun", size=10, bold=True)
    return row + 1


def write_block(
    ws,
    row: int,
    title: str,
    records: Sequence[Record],
    domain_key: str,
    backbone: str,
    include_w0_extras: bool,
    fixed_w: Optional[float] = None,
) -> int:
    group_fill = PatternFill("solid", fgColor="E2F0D9")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    cell = ws.cell(row=row, column=1, value=title)
    cell.fill = group_fill
    cell.font = Font(name="SimSun", size=10, bold=True)
    row += 1
    metric_rows: List[int] = []
    fixed_records = (
        best_common_fixed_weight_records(records, domain_key, backbone, fixed_w)
        if fixed_w is not None
        else {}
    )

    def add_record_row(
        variant: str,
        record: Optional[Record],
        label: str,
        note: str = "",
        is_reference: bool = False,
    ) -> None:
        nonlocal row
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value="pseudo-quality")
        if record:
            ws.cell(row=row, column=3, value=fmt_metric(record.recall10))
            ws.cell(row=row, column=4, value=fmt_metric(record.recall20))
            ws.cell(row=row, column=5, value=fmt_metric(record.ndcg10))
            ws.cell(row=row, column=6, value=fmt_metric(record.ndcg20))
            ws.cell(row=row, column=7, value=param_summary(record))
            ws.cell(row=row, column=8, value=str(record.path.relative_to(ROOT)))
        else:
            ws.cell(row=row, column=7, value=note)
        ws.cell(row=row, column=9, value=conclusion(backbone, variant, record, fixed=note if record else None))
        if is_reference:
            for col in range(1, 10):
                ws.cell(row=row, column=col).fill = REFERENCE_FILL
        else:
            metric_rows.append(row)
        row += 1

    for variant, _ in VARIANTS:
        if fixed_w is not None:
            rec = fixed_records.get(variant)
        else:
            candidates = filter_records(records, domain_key, backbone, variant)
            rec = best_record(candidates)
        label = short_variant(backbone, variant)
        if fixed_w is not None:
            note = f"pseudo_w={fixed_w:g}; shared setting"
        else:
            note = ""
        add_record_row(variant, rec, label, note=note)

        if include_w0_extras and variant in EXTRA_W0_VARIANTS:
            w0 = paired_weight_record(records, rec, pseudo_w=0.0)
            if w0:
                add_record_row(
                    variant,
                    w0,
                    short_variant(backbone, variant, pseudo_w0=True),
                    note="pseudo_w=0; same setting as above",
                    is_reference=True,
                )

    color_group_best(ws, metric_rows)
    return row + 1


def build_workbook(records: Sequence[Record], output: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "模型对比"
    row = 1

    for domain_key, title in DOMAINS:
        row = write_header(ws, row, title)
        for backbone in BACKBONES:
            row = write_block(
                ws,
                row,
                f"{backbone}：五类伪监督质量变体对比（含可用 pseudo_w=0 参考行）",
                records,
                domain_key,
                backbone,
                include_w0_extras=True,
                fixed_w=None,
            )
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    c = ws.cell(row=row, column=1, value="补充汇总：pseudo_w=0.1 固定权重下的五类变体对比")
    c.fill = PatternFill("solid", fgColor="FFF200")
    c.font = Font(name="SimSun", size=11, bold=True)
    row += 1

    for domain_key, title in DOMAINS:
        row = write_header(ws, row, title + "（pseudo_w=0.1）")
        for backbone in BACKBONES:
            row = write_block(
                ws,
                row,
                f"{backbone}：pseudo_w=0.1 固定权重汇总",
                records,
                domain_key,
                backbone,
                include_w0_extras=False,
                fixed_w=0.1,
            )
        row += 1

    apply_base_table_styles(ws)
    ws.freeze_panes = "A3"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def print_coverage(records: Sequence[Record]) -> None:
    print(f"[records] parsed={len(records)}")
    for domain_key, title in DOMAINS:
        print(f"\n[{title}]")
        for backbone in BACKBONES:
            for variant, _ in VARIANTS:
                all_records = filter_records(records, domain_key, backbone, variant)
                w01_records = filter_records(records, domain_key, backbone, variant, pseudo_w=0.1)
                w0_records = filter_records(records, domain_key, backbone, variant, pseudo_w=0.0)
                best = best_record(all_records)
                print(
                    f"  {backbone:<6} {variant:<18} "
                    f"all={len(all_records):>4} w0={len(w0_records):>3} w0.1={len(w01_records):>3} "
                    f"bestR20={best.recall20 if best else None}"
                )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output",
        default=str(ROOT / "experiment_records/unicdr_lsdcdr_pseudo_quality_records_20260623_augmented.xlsx"),
        help="Output workbook path.",
    )
    args = parser.parse_args()

    records: List[Record] = []
    for path in sorted(iter_log_paths()):
        records.extend(parse_log(path))

    print_coverage(records)
    output = Path(args.output)
    if not output.is_absolute():
        output = ROOT / output
    build_workbook(records, output)
    print(f"\n[OK] wrote {output}")


if __name__ == "__main__":
    main()
